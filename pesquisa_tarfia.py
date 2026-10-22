"""FareHub: pesquisa de tarifas em classe econômica."""

from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
import os
from pathlib import Path
import re
import sqlite3
import threading
from uuid import uuid4

from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
import requests


ROOT = Path(__file__).resolve().parent
DB_PATH = ROOT / "farehub.sqlite3"
ADVPS = (7, 30, 60)
MONTHLY_LIMITS = {"SerpApi": 250, "RapidAPI": 150}
NATIONAL = {"LATAM": "LATAM", "GOL": "GOL", "AZUL": "Azul"}
FEE_COLUMNS = {"taxa latam": "LATAM", "taxa gol": "GOL", "taxa azul": "Azul"}
SESSION = requests.Session()
LOCK = threading.Lock()
app = Flask(__name__)
app.json.ensure_ascii = False
app.config["TEMPLATES_AUTO_RELOAD"] = True
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0


def load_env_file():
    """Load .env file into os.environ if present, without overriding existing vars."""
    env_file = ROOT / ".env"
    if not env_file.exists():
        return
    for raw in env_file.read_text(encoding="utf-8-sig").splitlines():
        line = raw.strip()
        if "=" in line and not line.startswith("#"):
            k, v = line.split("=", 1)
            k, v = k.strip(), v.strip().strip('"\'')
            if k and k not in os.environ:
                os.environ[k] = v


def configured_apis():
    """Detect and return distinct, named APIs from environment variables and .env."""
    load_env_file()
    accounts = []
    seen_keys = set()

    # 1. Chaves enumeradas SerpApi (SERPAPI_KEY_1, SERPAPI_KEY_2...)
    serp_indices = sorted({
        int(m.group(1)) for k in os.environ
        if (m := re.match(r"^SERPAPI_KEY_(\d+)$", k, re.IGNORECASE))
    })
    for idx in serp_indices:
        key = os.environ.get(f"SERPAPI_KEY_{idx}", "").strip()
        name = os.environ.get(f"SERPAPI_NAME_{idx}", "").strip() or f"Conta {idx}"
        if key and key not in seen_keys:
            seen_keys.add(key)
            accounts.append({
                "id": f"SerpApi {idx} ({name})",
                "provider": "SerpApi",
                "account": name,
                "key": key
            })

    # 2. Chaves enumeradas RapidAPI (RAPIDAPI_KEY_1, RAPIDAPI_KEY_2...)
    rapid_indices = sorted({
        int(m.group(1)) for k in os.environ
        if (m := re.match(r"^RAPIDAPI_KEY_(\d+)$", k, re.IGNORECASE))
    })
    for idx in rapid_indices:
        key = os.environ.get(f"RAPIDAPI_KEY_{idx}", "").strip()
        name = os.environ.get(f"RAPIDAPI_NAME_{idx}", "").strip() or f"Conta {idx}"
        if key and key not in seen_keys:
            seen_keys.add(key)
            accounts.append({
                "id": f"RapidAPI {idx} ({name})",
                "provider": "RapidAPI",
                "account": name,
                "key": key
            })

    # 3. Chaves únicas simples (fallback)
    for key_name in ("SERPAPI_KEY", "CHAVE_API_SERPAPI"):
        val = os.environ.get(key_name, "").strip()
        if val and val not in seen_keys:
            seen_keys.add(val)
            accounts.append({
                "id": "SerpApi Principal",
                "provider": "SerpApi",
                "account": "Principal",
                "key": val
            })
            break

    for key_name in ("RAPIDAPI_KEY", "CHAVE_API_RAPIDAPI"):
        val = os.environ.get(key_name, "").strip()
        if val and val not in seen_keys:
            seen_keys.add(val)
            accounts.append({
                "id": "RapidAPI Principal",
                "provider": "RapidAPI",
                "account": "Principal",
                "key": val
            })
            break

    # 4. Formato legado do .env por seções
    if not accounts and (ROOT / ".env").exists():
        account = None
        for raw in (ROOT / ".env").read_text(encoding="utf-8-sig").splitlines():
            line = raw.strip()
            if line.upper().startswith("#CHAVE API'S EMAIL"):
                account = line.split("EMAIL", 1)[1].strip()
            elif line.startswith("#CHAVE API'S"):
                account = line.removeprefix("#CHAVE API'S").strip()
            elif "=" in line and not line.startswith("#"):
                key, value = line.split("=", 1)
                key, value = key.strip(), value.strip().strip('"\'')
                if key in ("CHAVE_API_SERPAPI", "CHAVE_API_RAPIDAPI") and value and value not in seen_keys:
                    seen_keys.add(value)
                    provider = "SerpApi" if key.endswith("SERPAPI") else "RapidAPI"
                    accounts.append({
                        "id": f"{provider} {len(accounts) + 1} ({account or 'Conta'})",
                        "provider": provider,
                        "account": account or "Conta",
                        "key": value
                    })

    return accounts


def connect_db():
    db = sqlite3.connect(DB_PATH)
    db.row_factory = sqlite3.Row
    db.execute("CREATE TABLE IF NOT EXISTS calls (id INTEGER PRIMARY KEY, api_id TEXT, provider TEXT, "
               "called_at TEXT DEFAULT CURRENT_TIMESTAMP, result TEXT, consumed INTEGER DEFAULT 1, "
               "remaining INTEGER, quota_limit INTEGER)")
    db.execute("CREATE TABLE IF NOT EXISTS exports (token TEXT PRIMARY KEY, rows_json TEXT NOT NULL, "
               "created_at TEXT DEFAULT CURRENT_TIMESTAMP)")
    db.commit()
    return db


def log_call(api, result, consumed=1, remaining=None, limit=None):
    with LOCK, connect_db() as db:
        db.execute("INSERT INTO calls (api_id, provider, result, consumed, remaining, quota_limit) "
                   "VALUES (?, ?, ?, ?, ?, ?)",
                   (api["id"], api["provider"], result, consumed, remaining, limit))
        db.commit()


def int_header(headers, key):
    try:
        return int(headers.get(key))
    except (TypeError, ValueError):
        return None


def rapid_quota(headers):
    """RapidAPI uses both billing headers and legacy request-quota headers."""
    remaining = (int_header(headers, "X-RapidAPI-Requests-Remaining")
                 if headers.get("X-RapidAPI-Requests-Remaining") is not None
                 else int_header(headers, "X-RateLimit-Requests-Remaining"))
    limit = (int_header(headers, "X-RapidAPI-Requests-Limit")
             if headers.get("X-RapidAPI-Requests-Limit") is not None
             else int_header(headers, "X-RateLimit-Requests-Limit"))
    return remaining, limit


def serp_balance(api):
    try:
        response = SESSION.get("https://serpapi.com/account.json",
                               params={"api_key": api["key"]}, timeout=12)
        response.raise_for_status()
        data = response.json()
        return {"used": data.get("this_month_usage"),
                "remaining": data.get("total_searches_left"),
                "limit": data.get("searches_per_month"), "source": "provedor"}
    except (requests.RequestException, ValueError):
        return {"used": None, "remaining": None, "limit": None,
                "source": "indisponível"}


def usage():
    with connect_db() as db:
        stats = {r["api_id"]: dict(r) for r in db.execute(
            "SELECT api_id, COUNT(*) AS attempts, COALESCE(SUM(consumed), 0) AS consumed "
            "FROM calls GROUP BY api_id")}
        rapid = {r["api_id"]: dict(r) for r in db.execute(
            "SELECT c.api_id, c.remaining, c.quota_limit FROM calls c JOIN "
            "(SELECT api_id, MAX(id) AS id FROM calls WHERE remaining IS NOT NULL "
            "GROUP BY api_id) x ON x.id = c.id")}
    rows = []
    for api in configured_apis():
        row = {k: api[k] for k in ("id", "provider")}
        row.update(stats.get(api["id"], {"attempts": 0, "consumed": 0}))
        if api["provider"] == "SerpApi":
            balance = serp_balance(api)
            if balance["limit"] is None:
                balance["limit"] = MONTHLY_LIMITS["SerpApi"]
                balance["source"] = "cota configurada; saldo indisponível"
            row.update(balance)
        else:
            latest = rapid.get(api["id"], {})
            remaining = latest.get("remaining")
            limit = latest.get("quota_limit") or MONTHLY_LIMITS["RapidAPI"]
            row.update({"remaining": remaining, "limit": limit,
                        "used": limit - remaining if remaining is not None and 0 <= remaining <= limit else None,
                        "source": "cabeçalho da última chamada" if remaining is not None else "cota configurada; saldo indisponível"})
        rows.append(row)
    return rows


def serp_search(api, origin, destination, departure):
    params = {"engine": "google_flights", "departure_id": origin,
              "arrival_id": destination, "outbound_date": departure,
              "type": 2, "travel_class": 1, "currency": "BRL", "hl": "pt",
              "gl": "br", "api_key": api["key"]}
    try:
        response = SESSION.get("https://serpapi.com/search.json", params=params, timeout=35)
        data = response.json()
        metadata = data.get("search_metadata", {})
        # SerpApi does not charge cached searches.
        consumed = 0 if metadata.get("status") == "Cached" else 1
        error = data.get("error")
        if response.status_code >= 400 or error:
            log_call(api, "erro", consumed)
            raise RuntimeError(str(error or f"HTTP {response.status_code}"))
        log_call(api, "sucesso", consumed)
        return data
    except (requests.RequestException, ValueError) as exc:
        log_call(api, "erro")
        raise RuntimeError(type(exc).__name__) from exc


def rapid_search(api, origin, destination, departure):
    # The RapidAPI subscription/host is separate from the application ID in .env.
    host = os.environ.get("RAPIDAPI_HOST", "google-flights2.p.rapidapi.com")
    path = os.environ.get("RAPIDAPI_SEARCH_PATH", "/api/v1/searchFlights")
    if not re.fullmatch(r"[a-z0-9.-]+\.p\.rapidapi\.com", host):
        raise RuntimeError("RAPIDAPI_HOST inválido")
    try:
        response = requests.get(f"https://{host}{path}",
            headers={"X-RapidAPI-Key": api["key"], "X-RapidAPI-Host": host},
            params={"departure_id": origin, "arrival_id": destination,
                    "outbound_date": departure, "travel_class": "ECONOMY",
                    "currency": "BRL", "language_code": "pt-BR", "country_code": "BR",
                    "search_type": "cheap", "adults": "1"}, timeout=40)
        remaining, limit = rapid_quota(response.headers)
        data = response.json()
        if response.status_code >= 400 or data.get("status") is False:
            log_call(api, "erro", 1, remaining, limit)
            raise RuntimeError(f"HTTP {response.status_code}")
        log_call(api, "sucesso", 1, remaining, limit)
        return data
    except (requests.RequestException, ValueError) as exc:
        log_call(api, "erro")
        raise RuntimeError(type(exc).__name__) from exc


def unwrap_flights(data):
    """Accept SerpApi and common Google Flights RapidAPI wrappers."""
    current = data
    for _ in range(4):
        if not isinstance(current, dict):
            break
        if "best_flights" in current or "other_flights" in current:
            return list(current.get("best_flights") or []) + list(current.get("other_flights") or [])
        if "topFlights" in current or "otherFlights" in current:
            return list(current.get("topFlights") or []) + list(current.get("otherFlights") or [])
        current = current.get("data") or current.get("itineraries") or {}
    return []


def airline_name(leg):
    name = str(leg.get("airline") or "").strip()
    number = str(leg.get("flight_number") or "").upper().replace(" ", "")
    code = re.match(r"^[A-Z0-9]{2}", number)
    if code and code.group() in ("LA", "JJ") or "LATAM" in name.upper():
        return "LATAM"
    if code and code.group() == "G3" or name.upper().startswith("GOL"):
        return "GOL"
    if code and code.group() == "AD" or name.upper().startswith("AZUL"):
        return "Azul"
    return name or (code.group() if code else "Não informada")


def quoted_component(item, *names):
    """Use a fare component only when the provider returns it explicitly as money."""
    for values in (item, item.get("price_breakdown")):
        if not isinstance(values, dict):
            continue
        for name in names:
            value = values.get(name)
            if isinstance(value, (int, float)) and not isinstance(value, bool) and value >= 0:
                return round(float(value), 2)
    return None


def normalize(data, origin, destination, departure, advp, scope, source):
    best = {}
    for item in unwrap_flights(data):
        legs = item.get("flights") or item.get("legs") or []
        if not legs and item.get("airlines"):
            airlines = item["airlines"]
            if not isinstance(airlines, list) or len(set(airlines)) != 1:
                continue
            legs = [{"airline": airlines[0], "flight_number": item.get("flight_number"),
                     "departure_airport": {"time": item.get("departure_time")},
                     "arrival_airport": {"id": destination}}]
        if not legs or not isinstance(legs, list):
            continue
        names = {airline_name(leg) for leg in legs}
        if len(names) != 1:  # No company can own a mixed-carrier fare.
            continue
        airline = names.pop()
        if scope == "nacional" and airline not in NATIONAL.values():
            continue
        price_raw = item.get("price")
        try:
            price = float(price_raw)
        except (ValueError, TypeError):
            continue
        if price < 0:
            continue
        first = legs[0]
        dep = first.get("departure_airport") or {}
        time = str(dep.get("time") or "")
        match = re.search(r"(\d{2}:\d{2})", time)
        route = " → ".join([origin] + [str((leg.get("arrival_airport") or {}).get("id") or
                                         (leg.get("arrival_airport") or {}).get("code") or
                                         (leg.get("arrival_airport") or {}).get("airport_code") or destination)
                                   for leg in legs])
        row = {"advp": advp, "data": departure, "origem": origin,
               "destino": destination, "cia": airline,
               "voo": " + ".join(str(leg.get("flight_number") or "—") for leg in legs),
               "hora": match.group(1) if match else "—", "trecho": route,
               "tarifa": quoted_component(item, "base_fare", "baseFare"),
               "taxa": quoted_component(item, "taxes", "tax_amount", "taxAmount"),
               "desconto": quoted_component(item, "discount_amount", "discountAmount"),
               "preco_total": price,
               "total_aproximado": isinstance(price_raw, int) and not isinstance(price_raw, bool),
               "classe": "Economy", "moeda": "BRL", "fonte": source}
        if airline not in best or price < best[airline]["preco_total"]:
            best[airline] = row
    return sorted(best.values(), key=lambda row: row["preco_total"])


def search_one(origin, destination, departure, advp, scope, api_choice="automatico"):
    errors = []
    apis = configured_apis()
    ordered = ([a for a in apis if a["provider"] == "SerpApi"] +
               [a for a in apis if a["provider"] == "RapidAPI"])
    if api_choice != "automatico":
        ordered = [a for a in ordered if a["id"] == api_choice]
    for api in ordered:
        try:
            data = (serp_search if api["provider"] == "SerpApi" else rapid_search)(
                api, origin, destination, departure)
            if not unwrap_flights(data):
                # A valid empty answer should not spend more accounts.
                return [], None
            return normalize(data, origin, destination, departure, advp, scope,
                             f'{api["id"]} · {api["provider"]}'), None
        except RuntimeError as exc:
            errors.append(f'{api["id"]}: {exc}')
    return [], "Nenhuma API retornou esta busca. " + "; ".join(errors)


def parse_routes(file):
    if not file or not file.filename.lower().endswith(".xlsx"):
        raise ValueError("Envie uma planilha .xlsx.")
    try:
        workbook = load_workbook(file.stream, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header = [re.sub(r"[\s_]+", " ", str(v or "").strip().casefold())
                  for v in next(iterator)]
        if "origem" not in header or "destino" not in header:
            raise ValueError("A primeira linha deve conter as colunas origem e destino.")
        ix_o, ix_d = header.index("origem"), header.index("destino")
        fee_indexes = {airline: header.index(label) if label in header else None
                       for label, airline in FEE_COLUMNS.items()}
        routes = []
        for line_number, row in enumerate(iterator, start=2):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            origin = str(row[ix_o] or "").strip().upper() if ix_o < len(row) else ""
            dest = str(row[ix_d] or "").strip().upper() if ix_d < len(row) else ""
            if not valid_iata(origin) or not valid_iata(dest) or origin == dest:
                raise ValueError(f"Trecho inválido na linha {line_number}. Use códigos IATA de 3 letras.")
            fees = {airline: parse_fee(row[index] if index is not None and index < len(row) else None,
                                       label, line_number)
                    for label, airline in FEE_COLUMNS.items()
                    for index in [fee_indexes[airline]]}
            routes.append({"origin": origin, "destination": dest, "fees": fees})
            if len(routes) > 10:
                raise ValueError("A planilha aceita até 10 trechos.")
        workbook.close()
        if not routes:
            raise ValueError("A planilha não contém trechos.")
        unique = {}
        for route in routes:
            key = (route["origin"], route["destination"],
                   *(route["fees"][airline] for airline in FEE_COLUMNS.values()))
            unique.setdefault(key, route)
        return list(unique.values())
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Não foi possível ler o arquivo XLSX.") from exc


def valid_iata(value):
    return bool(re.fullmatch(r"[A-Z]{3}", value))


def parse_fee(value, label, line_number):
    if value is None or isinstance(value, str) and not value.strip():
        return Decimal("0.00")
    try:
        text = str(value).strip().replace("R$", "").replace(" ", "")
        if "," in text:
            text = text.replace(".", "").replace(",", ".")
        amount = Decimal(text)
        if not amount.is_finite() or amount < 0 or amount.as_tuple().exponent < -2:
            raise InvalidOperation
        return amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{label.title()} inválida na linha {line_number}. Use valor em BRL, por exemplo 33,44.") from exc


def make_xlsx(rows, example=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Exemplo" if example else "Tarifas"
    headers = ["origem", "destino", "taxa latam", "taxa gol", "taxa azul"] if example else [
        "ADVP", "Data da viagem", "Origem", "Destino", "Cia", "Número do voo",
        "Hora do voo", "Trecho", "Tarifa consultada: total (BRL)",
        "Taxa (BRL)", "Tarifa net (BRL)", "Total sem centavos", "Classe", "Fonte"]
    ws.append(headers)
    if example:
        ws.append(["CGB", "GRU", 0, 0, 0])
        ws.append(["GRU", "GIG", None, None, None])
    else:
        for row in rows:
            legacy = "preco_total" not in row
            values = [row.get(k) for k in ("advp", "data", "origem", "destino",
                     "cia", "voo", "hora", "trecho")]
            total = row.get("preco_total", row.get("tarifa"))
            fee = row.get("taxa_embarque", 0)
            net = row.get("tarifa_net")
            if net is None and isinstance(total, (int, float)):
                net = round(float(total) - float(fee), 2)
            values.extend([total, fee, net,
                           "Sim" if row.get("total_aproximado", legacy) else "Não",
                           row.get("classe"), row.get("fonte")])
            ws.append([value if value is not None else "Não informado" for value in values])
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="087CBC")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = min(38, max(13, max(len(str(c.value or "")) for c in col) + 3))
    if not example:
        for letter in ("I", "J", "K"):
            for cell in ws[letter][1:]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '"R$" #,##0.00'
    else:
        for letter in ("C", "D", "E"):
            for cell in ws[letter][1:]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '"R$" #,##0.00'
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


@app.get("/")
def home():
    return render_template("index.html", today=date.today().isoformat(), apis=configured_apis())


@app.get("/api/usage")
def api_usage():
    return jsonify(usage())


@app.post("/api/usage/refresh-rapid")
def refresh_rapid_usage():
    rapid_apis = [a for a in configured_apis() if a["provider"] == "RapidAPI"]
    departure = (date.today() + timedelta(days=7)).isoformat()

    def probe(api):
        try:
            rapid_search(api, "CGB", "GRU", departure)
            return None
        except RuntimeError as exc:
            return f'{api["id"]}: {exc}'

    with ThreadPoolExecutor(max_workers=3) as pool:
        errors = [error for error in pool.map(probe, rapid_apis) if error]
    return jsonify(apis=usage(), errors=errors, probed=len(rapid_apis))


@app.get("/exemplo.xlsx")
def example_xlsx():
    return send_file(make_xlsx([], True), as_attachment=True,
                     download_name="exemplo_trechos.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.post("/api/search")
def api_search():
    today = date.today()
    mode = request.form.get("mode", "single")
    if mode not in ("single", "data_especifica", "planilha"):
        return jsonify(error="Escolha um tipo de pesquisa válido."), 400
    scope = request.form.get("scope", "nacional")
    if scope not in ("nacional", "internacional"):
        return jsonify(error="Escolha cias nacionais ou todas as cias."), 400
    api_choice = request.form.get("api_choice", "automatico")
    if api_choice != "automatico" and api_choice not in {a["id"] for a in configured_apis()}:
        return jsonify(error="Escolha uma API configurada."), 400
    try:
        if mode == "planilha":
            routes = parse_routes(request.files.get("file"))
        else:
            origin = request.form.get("origin", "").strip().upper()
            destination = request.form.get("destination", "").strip().upper()
            if not valid_iata(origin) or not valid_iata(destination) or origin == destination:
                raise ValueError("Informe origem e destino com códigos IATA de 3 letras diferentes.")
            routes = [{"origin": origin, "destination": destination, "fees": {}}]
        if mode == "data_especifica":
            departure_text = request.form.get("departure_date", "").strip()
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", departure_text):
                raise ValueError("Informe uma data de embarque válida.")
            try:
                chosen_date = date.fromisoformat(departure_text)
            except ValueError as exc:
                raise ValueError("Informe uma data de embarque válida.") from exc
            if chosen_date < today:
                raise ValueError("A data de embarque não pode estar no passado.")
    except ValueError as exc:
        return jsonify(error=str(exc)), 400
    rows, errors = [], []
    departures = ([(departure_text, (chosen_date - today).days)]
                  if mode == "data_especifica" else
                  [((today + timedelta(days=advp)).isoformat(), advp)
                   for advp in ADVPS])
    for route in routes:
        origin, destination = route["origin"], route["destination"]
        for departure, advp in departures:
            found, error = search_one(origin, destination, departure, advp, scope, api_choice)
            for row in found:
                fee = route["fees"].get(row["cia"], Decimal("0.00"))
                total = Decimal(str(row["preco_total"]))
                row["taxa_embarque"] = float(fee)
                row["tarifa_net"] = float((total - fee).quantize(Decimal("0.01")))
            rows.extend(found)
            if error:
                errors.append(f"{origin} → {destination}, {advp} dias: {error}")
    token = uuid4().hex
    with connect_db() as db:
        db.execute("INSERT INTO exports (token, rows_json) VALUES (?, ?)",
                   (token, json.dumps(rows, ensure_ascii=False)))
        db.commit()
    return jsonify(rows=rows, errors=errors, token=token, routes=len(routes),
                   searches=len(routes) * len(departures))


@app.get("/export/<token>.xlsx")
def export(token):
    if not re.fullmatch(r"[a-f0-9]{32}", token):
        return "Arquivo não encontrado", 404
    with connect_db() as db:
        result = db.execute("SELECT rows_json FROM exports WHERE token = ?", (token,)).fetchone()
    if result is None:
        return "Arquivo não encontrado", 404
    return send_file(make_xlsx(json.loads(result["rows_json"])), as_attachment=True,
                     download_name="tarifas_farehub.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8502))
    app.run(host="0.0.0.0", port=port, debug=False)
